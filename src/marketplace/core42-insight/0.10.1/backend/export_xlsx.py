"""Insight revenue as an Excel workbook, in the look of Azure Cost Manager's
export (the operator's own workbook): Office theme, Aptos Narrow, no gridlines,
navy section bars, #D9E1F2 table headers, thin grey borders, `"$"#,##0` money.
Values and a chart only -- no formulas.

`openpyxl` is provided to plugin backends by the host (framework API 9).
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .office_theme import THEME_XML

FONT = "Aptos Narrow"
NAVY, BLUE_TAB, GREY_TEXT = "1F3864", "2E75B6", "595959"
HEADER_FILL, FORECAST_FILL, GRID = "D9E1F2", "F2F2F2", "BFBFBF"
MONEY = '"$"#,##0;\\("$"#,##0\\);\\-'
PERCENT = '0.0%;\\(0.0%\\);\\-'
MONTH = 'yyyy\\-mm'

THIN = Side(style="thin", color=GRID)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BOX = Border(left=THIN, right=THIN, top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY))
MODEL_NAMES = {"percentage": "% of consumption", "upfront": "Up front", "yearly": "Yearly"}


def _font(size: float = 11, bold: bool = False, italic: bool = False, color: str | None = None) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _month_date(month: str) -> datetime:
    return datetime(int(month[:4]), int(month[5:]), 1)


def _workbook() -> Workbook:
    workbook = Workbook()
    workbook.loaded_theme = THEME_XML.encode("utf8")
    return workbook


def _sheet_title(name: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "Customer"
    title, number = base, 2
    while title.lower() in taken:
        title, number = f"{base[:25]} {number}", number + 1
    taken.add(title.lower())
    return title


def _new_sheet(workbook: Workbook, title: str, heading: str, subtitle: str, sheet=None):
    sheet = sheet or workbook.create_sheet(title)
    sheet.title = title
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = BLUE_TAB
    sheet["A1"], sheet["A2"] = heading, subtitle
    sheet["A1"].font = _font(16, bold=True)
    sheet["A2"].font = _font(italic=True, color=GREY_TEXT)
    sheet.row_dimensions[1].height = 23.5
    return sheet


def _section(sheet, row: int, title: str, last_column: int) -> None:
    for column in range(1, last_column + 1):
        sheet.cell(row, column).fill = _fill(NAVY)
    sheet.cell(row, 1, title).font = _font(bold=True, color="FFFFFF")


def _table(sheet, top: int, headers: list[str], rows: list[list], formats: dict[int, str],
           shaded: set[int] | None = None, total_row: list | None = None) -> int:
    for offset, header in enumerate(headers):
        cell = sheet.cell(top, 1 + offset, header)
        cell.font, cell.fill, cell.border = _font(bold=True), _fill(HEADER_FILL), BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    every = rows + ([total_row] if total_row else [])
    for index, values in enumerate(every, start=1):
        is_total = total_row is not None and index == len(every)
        for offset, value in enumerate(values):
            cell = sheet.cell(top + index, 1 + offset, value)
            cell.font = _font(bold=is_total, italic=bool(shaded and index - 1 in shaded))
            cell.border = TOTAL_BOX if is_total else BOX
            if shaded and index - 1 in shaded:
                cell.fill = _fill(FORECAST_FILL)
            if offset in formats and value is not None and not isinstance(value, str):
                cell.number_format = formats[offset]
    return top + len(every)


def _widths(sheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _kpis(sheet, top: int, summary: dict) -> int:
    _section(sheet, top, f"{summary['year']} at a glance", 6)
    last = summary.get("last_month")
    rows = [
        ["Last closed month", last["month"] if last else "-", last["revenue"] if last else None],
        ["Year to date", summary["year"], summary["ytd"]],
        ["Forecast to year end", summary["year"], summary["forecast_year_end"]],
        ["Lifetime (actual)", "", summary["lifetime"]],
    ]
    return _table(sheet, top + 1, ["Measure", "Period", "USD"], rows, {2: MONEY})


def _monthly(sheet, top: int, rows: list[dict], last_column: int = 6) -> int:
    _section(sheet, top, "Month by month (forecast months shaded)", last_column)
    table = [[_month_date(row["month"]), "Actual" if row["kind"] == "actual" else "Forecast",
              row.get("line") or "", row["consumption"], row["revenue"], row["cumulative"]] for row in rows]
    shaded = {index for index, row in enumerate(rows) if row["kind"] != "actual"}
    end = _table(sheet, top + 1, ["Month", "Status", "Terms", "Consumption", "Insight revenue", "Cumulative"],
                 table, {0: MONTH, 3: MONEY, 4: MONEY, 5: MONEY}, shaded,
                 total_row=["Total", "", "", None, round(sum(row["revenue"] for row in rows), 2), None])
    if rows:
        chart = BarChart()
        chart.title, chart.height, chart.width = "Insight revenue by month", 7.5, 18
        chart.add_data(Reference(sheet, min_col=5, min_row=top + 1, max_row=top + len(rows) + 1), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=top + 2, max_row=top + len(rows) + 1))
        chart.y_axis.numFmt = MONEY
        chart.x_axis.number_format = MONTH
        line = LineChart()
        line.add_data(Reference(sheet, min_col=6, min_row=top + 1, max_row=top + len(rows) + 1), titles_from_data=True)
        line.y_axis.axId = 200
        line.y_axis.crosses = "max"
        chart += line
        sheet.add_chart(chart, f"H{top + 1}")
    return end


def _customer_sheet(sheet, report: dict, today: date) -> None:
    _new_sheet(sheet.parent, sheet.title, f"Core42 Insight — {report['name']}",
               f"USD. Generated {today.isoformat()}; months to {report['current_month']} exclusive are actual.", sheet)
    row = _kpis(sheet, 4, report["summary"]) + 2
    _section(sheet, row, "Terms", 6)
    term_rows = []
    for member in report["members"]:
        for line in member["terms"]:
            rate = line.get("percent") / 100 if line["model"] == "percentage" else line.get("amount")
            basis = ""
            if line["model"] == "percentage":
                basis = ("incl. Marketplace" if line["include_marketplace"] else "excl. Marketplace") + ", " + \
                        ("less credits" if line["subtract_credits"] else "credits not subtracted")
            if line.get("payments"):
                basis = "; ".join(f"{payment['date']} ${payment['amount']:,.0f}" for payment in line["payments"])
            term_rows.append([member["name"], MODEL_NAMES[line["model"]], line["start"], line.get("end") or "open",
                              rate, basis])
    terms_top = row + 1
    row = _table(sheet, terms_top, ["Customer", "Model", "From", "To", "Rate / amount", "Basis"],
                 term_rows or [["No terms", "", "", "", None, ""]], {}) + 2
    for index, values in enumerate(term_rows):
        sheet.cell(terms_top + 1 + index, 5).number_format = \
            PERCENT if values[1] == MODEL_NAMES["percentage"] else MONEY
    _monthly(sheet, row, report["rows"])
    _widths(sheet, {"A": 26, "B": 18, "C": 34, "D": 14, "E": 16, "F": 16})


def customer_workbook(report: dict, today: date | None = None) -> tuple[str, bytes]:
    today = today or date.today()
    workbook = _workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(report["name"], set())
    _customer_sheet(sheet, report, today)
    return f"Core42 Insight - {report['name']} - {today.isoformat()}.xlsx", _saved(workbook)


def portfolio_workbook(portfolio: dict, reports: list[dict], today: date | None = None) -> tuple[str, bytes]:
    """The across-customers sheet first, then one sheet per top-level customer
    (its affiliates rolled up in it)."""
    today = today or date.today()
    workbook = _workbook()
    sheet = _new_sheet(workbook, "All customers", "Core42 Insight — all customers",
                       f"USD. Generated {today.isoformat()}. Each customer's own figures; affiliates indented.",
                       workbook.active)
    row = _kpis(sheet, 4, portfolio["totals"]) + 2
    _section(sheet, row, "Customers", 6)
    table = [[("    " if customer["parent"] else "") + customer["name"], customer.get("in_force") or "none today",
              (customer["summary"]["last_month"] or {}).get("revenue"), customer["summary"]["ytd"],
              customer["summary"]["forecast_year_end"], customer["summary"]["lifetime"]]
             for customer in portfolio["customers"]]
    totals = portfolio["totals"]
    row = _table(sheet, row + 1, ["Customer", "Terms today", "Last month", "YTD", "Forecast to year end", "Lifetime"],
                 table, {2: MONEY, 3: MONEY, 4: MONEY, 5: MONEY},
                 total_row=["Total", "", (totals["last_month"] or {}).get("revenue"), totals["ytd"],
                            totals["forecast_year_end"], totals["lifetime"]]) + 2
    _monthly(sheet, row, portfolio["months"])
    _widths(sheet, {"A": 30, "B": 30, "C": 14, "D": 14, "E": 18, "F": 14})
    taken = {"all customers"}
    for report in reports:
        _customer_sheet(workbook.create_sheet(_sheet_title(report["name"], taken)), report, today)
    return f"Core42 Insight - all customers - {today.isoformat()}.xlsx", _saved(workbook)


def _saved(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
