"""Reads one captured attachment as plain text, for summarization.

Why this exists: the agent's own `read_file` cannot open Office binaries, and
the first run of this Skill wrote a "no extractable text; use filename context
only" placeholder for 36 of 71 files because of it. Pointing the agent at hub
skills per extension did not fix that -- a scheduled agent cannot reliably
reach them. So extraction is done HERE, in code, the same way `read_thread.py`
strips email HTML: at read, never at capture. The original file stays exactly
as it arrived, so a summary that looks wrong can be checked against it.

    python read_file.py --file-note "<Work/Threads/<T>/files/<slug>/<note>.md>"
    python read_file.py --file-note ... --max-chars 40000

Prints a header (filename, type, size, source Thread) and then the extracted
text. It never decides what the file MEANS -- that is the reader's job -- but
it does say what it could not do: no text layer, encrypted, truncated,
unsupported. An honest "could not read this" is worth more than a confident
summary of a filename.

Needs pypdf, python-docx, openpyxl and python-pptx. They live in a dedicated
venv, not in the uv-managed Hermes runtime, which refuses modification.
"""
from __future__ import annotations

import argparse
import email
import html
import os
import re
import sys
import zipfile
from email import policy
from pathlib import Path

import vault_manager as vm

# Bounded by CONTEXT, not by the document: at ~4 chars a token this is ~10k
# tokens, so a batch of ten files stays near 100k tokens of reading.
_DEFAULT_MAX_CHARS = 40_000
# A spreadsheet is the format most likely to be enormous and least likely to
# need every row to be understood.
_SHEET_ROW_CAP = 200
# Below this, an image attached to an email is almost always a signature logo
# or an inline icon. Said as a hint -- the reader still decides.
_SMALL_IMAGE_BYTES = 20_000

_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}
_TEXT_EXT = {".txt", ".csv", ".md", ".json", ".log", ".xml", ".tsv"}
_HTML_EXT = {".html", ".htm"}

_DROP_WHOLE = re.compile(r"<(script|style|head|title)\b[^>]*>.*?</\1>", re.I | re.S)
_BLOCK_BREAK = re.compile(r"</?(p|div|br|tr|li|h[1-6]|table|blockquote)\b[^>]*>", re.I)
_ANY_TAG = re.compile(r"<[^>]+>")
_MANY_BLANKS = re.compile(r"\n{3,}")
_TRAILING_SPACE = re.compile(r"[ \t]+\n")


def html_to_text(raw: str) -> str:
    if not raw:
        return ""
    if "<" not in raw:
        return raw.strip()
    text = _DROP_WHOLE.sub(" ", raw)
    text = _BLOCK_BREAK.sub("\n", text)
    text = _ANY_TAG.sub("", text)
    text = html.unescape(text)
    text = _TRAILING_SPACE.sub("\n", text)
    return _MANY_BLANKS.sub("\n\n", text).strip()


def _pdf(path: Path) -> tuple[str, list[str]]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return "", ["encrypted PDF -- no text available without its password"]
    parts, empty = [], 0
    for number, page in enumerate(reader.pages, 1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if not text:
            empty += 1
            continue
        parts.append(f"--- page {number} ---\n{text}")
    notes = [f"{len(reader.pages)} pages"]
    if reader.pages and empty == len(reader.pages):
        notes.append("no text layer on any page -- likely scanned or image-only; "
                     "summarize from the filename and the Thread, and SAY that is what you did")
    elif empty:
        notes.append(f"{empty} page(s) had no extractable text")
    return "\n\n".join(parts), notes


def _docx(path: Path) -> tuple[str, list[str]]:
    import docx
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for number, table in enumerate(document.tables, 1):
        rows = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                text = cell.text.strip()
                # python-docx repeats a merged cell once per grid column it spans.
                if not cells or cells[-1] != text:
                    cells.append(text)
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"--- table {number} ---\n" + "\n".join(rows))
    return "\n".join(parts), []


def _xlsx(path: Path) -> tuple[str, list[str]]:
    import openpyxl
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts, notes = [], []
    try:
        for sheet in workbook.worksheets:
            rows, total = [], 0
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not any(cells):
                    continue
                total += 1
                if total <= _SHEET_ROW_CAP:
                    rows.append(" | ".join(cells))
            if total > _SHEET_ROW_CAP:
                notes.append(f"sheet {sheet.title!r}: first {_SHEET_ROW_CAP} of {total} "
                             "non-empty rows shown")
            header = f"--- sheet: {sheet.title} ({total} rows) ---"
            parts.append(header + "\n" + "\n".join(rows) if rows else header)
    finally:
        workbook.close()
    return "\n\n".join(parts), notes


def _shape_texts(shapes) -> list[str]:
    """Text from every shape, descending into groups -- a grouped text box is
    common in decks and invisible to a flat walk."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    lines: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_shape_texts(shape.shapes))
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                lines.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    return lines


def _pptx(path: Path) -> tuple[str, list[str]]:
    from pptx import Presentation
    presentation = Presentation(str(path))
    parts = []
    for number, slide in enumerate(presentation.slides, 1):
        lines = _shape_texts(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            speaker = slide.notes_slide.notes_text_frame.text.strip()
            if speaker:
                lines.append(f"[speaker notes] {speaker}")
        if lines:
            parts.append(f"--- slide {number} ---\n" + "\n".join(lines))
    slide_count = len(presentation.slides)
    notes = [f"{slide_count} slides"]
    if slide_count and not parts:
        notes.append("no text on any slide -- likely image-only slides")
    return "\n\n".join(parts), notes


def _eml(path: Path) -> tuple[str, list[str]]:
    message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    head = [f"{name}: {message[name]}" for name in ("From", "To", "Date", "Subject")
            if message[name]]
    body = message.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        content = body.get_content()
        text = html_to_text(content) if body.get_content_type() == "text/html" else content.strip()
    inner = [part.get_filename() for part in message.iter_attachments() if part.get_filename()]
    notes = ["a forwarded email"]
    if inner:
        notes.append("attachments inside it: " + ", ".join(inner))
    return "\n".join(head) + "\n\n" + text, notes


def _ics(path: Path) -> tuple[str, list[str]]:
    keep = ("SUMMARY", "DTSTART", "DTEND", "LOCATION", "ORGANIZER", "DESCRIPTION")
    lines = [line for line in path.read_text(encoding="utf-8", errors="replace").splitlines()
             if line.split(":", 1)[0].split(";", 1)[0] in keep]
    return "\n".join(lines), ["a calendar invite"]


def _zip(path: Path) -> tuple[str, list[str]]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    return "\n".join(names[:200]), [f"an archive of {len(names)} entries -- contents "
                                    "LISTED, not extracted"]


def _image(path: Path) -> tuple[str, list[str]]:
    size = path.stat().st_size
    notes = [f"an image ({size // 1024} KB) -- view it at the PATH above if you can read "
             "images; otherwise summarize from the filename and the Thread"]
    if size < _SMALL_IMAGE_BYTES:
        notes.append("small enough that in an email it is almost always a signature logo "
                     "or an inline icon -- one honest line is a complete summary")
    return "", notes


def _text(path: Path) -> tuple[str, list[str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    return (html_to_text(raw) if path.suffix.lower() in _HTML_EXT else raw.strip()), []


_EXTRACTORS = {
    ".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx,
    ".eml": _eml, ".ics": _ics, ".zip": _zip,
}


def extract(path: Path) -> tuple[str, list[str]]:
    """Never raises. A corrupt or unexpected file is reported, not crashed on:
    one bad attachment must not stop a batch, and the reader can still write an
    honest line about a file it could not open."""
    suffix = path.suffix.lower()
    if suffix in _IMAGE_EXT:
        handler = _image
    elif suffix in _TEXT_EXT or suffix in _HTML_EXT:
        handler = _text
    else:
        handler = _EXTRACTORS.get(suffix)
    if handler is None:
        return "", [f"unsupported format {suffix or '(no extension)'} -- summarize from "
                    "the filename and the Thread"]
    try:
        return handler(path)
    except Exception as exc:
        return "", [f"extraction failed ({type(exc).__name__}: {str(exc)[:160]}) -- the file "
                    "may be corrupt or password-protected; say so rather than guess"]


def _original_file(note: Path, frontmatter: dict) -> Path | None:
    """Scanned through long_path. Past Windows' 260-character MAX_PATH a plain
    iterdir/is_file silently reports nothing, which here would read as "the
    file was never captured" -- a false statement the reader then repeats in
    its summary."""
    with os.scandir(vm.long_path(note.parent)) as entries:
        names = [e.name for e in entries if e.is_file() and not e.name.lower().endswith(".md")]
    wanted = (frontmatter.get("original_filename") or "").strip()
    if wanted in names:
        return note.parent / wanted
    return note.parent / names[0] if len(names) == 1 else None


def read_file(note: Path, *, max_chars: int = _DEFAULT_MAX_CHARS) -> str:
    frontmatter, _ = vm.read_note(note)
    original = _original_file(note, frontmatter)
    name = frontmatter.get("original_filename") or (original.name if original else note.stem)

    header = [f"FILE: {name}", f"NOTE: {note}"]
    if frontmatter.get("source_thread"):
        header.append(f"THREAD: {frontmatter['source_thread']}")

    if original is None:
        header.append("NOTES: the original file is not in the vault -- capture may have "
                      "skipped it (size cap). Summarize from the filename and the Thread, "
                      "and say the file itself was not available.")
        return "\n".join(header) + "\n--- content ---\n(no file)"

    # Opened through the long-path form; the plain path is what gets SHOWN.
    readable = Path(vm.long_path(original))
    text, notes = extract(readable)
    header.insert(1, f"TYPE: {original.suffix.lower() or '(none)'} | "
                     f"SIZE: {readable.stat().st_size // 1024} KB")
    header.insert(2, f"PATH: {original}")
    if notes:
        header.append("NOTES: " + "; ".join(notes))

    body = text.strip() or "(no extractable text)"
    if max_chars > 0 and len(body) > max_chars:
        shown = body[:max_chars]
        body = (shown + f"\n\n[TRUNCATED: showed {max_chars:,} of {len(text):,} characters -- "
                "your summary covers only the start of this file; say so]")
    return "\n".join(header) + "\n--- content ---\n" + body


def main() -> int:
    parser = argparse.ArgumentParser(description="Read one captured attachment as text.")
    parser.add_argument("--file-note", required=True,
                        help="The attachment's companion .md note (not the file itself).")
    parser.add_argument("--max-chars", type=int, default=_DEFAULT_MAX_CHARS,
                        help="Truncate the extracted text beyond this. 0 = no limit.")
    args = parser.parse_args()
    note = Path(args.file_note)
    if not os.path.isfile(vm.long_path(note)):
        print(f"no File note at {note}")
        return 2
    sys.stdout.reconfigure(encoding="utf-8")
    print(read_file(note, max_chars=args.max_chars))
    return 0


if __name__ == "__main__":
    sys.exit(main())
